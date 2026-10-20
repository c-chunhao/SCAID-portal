"""Source-recorded B-state, individual-level IFN and matched-population sensitivity.
No cells are replicates; no published main estimate is overwritten.
"""
from pathlib import Path
import os
import hashlib,json,datetime,sys
import numpy as np,pandas as pd,statsmodels.formula.api as smf
from statsmodels.stats.multitest import multipletests
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
W=Path(os.environ.get('SCAID_BCELL_OUTPUT', 'work/bcell_state')).resolve();P=W/'provenance';T=W/'tables';OLD=Path(os.environ.get('SCAID_BASELINE_SCIENCE', 'data/baseline_science')).resolve()
plan=json.loads((P/'analysis_plan.json').read_text());inp=json.loads((P/'input_manifest.json').read_text());scan=json.loads((P/'selected_B_scan.json').read_text());states=plan['candidate_states'];watch=json.loads((P/'watch_genes.json').read_text());IFN=plan['IFN_genes']
sha=lambda p:hashlib.file_digest(open(p,'rb'),'sha256').hexdigest()
assert sha(P/'analysis_plan.json')==inp['analysis_plan_sha256']
assert scan['entries_read']==1011675483 and scan['source_cells']==1263676 and not scan['whole_matrix_materialized']
b=pd.read_csv(T/'selected_B_cells_source_mapping.tsv.gz',sep='\t');per=pd.read_csv(T/'B_state_counts_by_individual.tsv',sep='\t').set_index('unit_id');assert per.ind_cov.is_unique
lib=np.fromfile(P/'selected_B_library.uint64',dtype='<u8');counts=np.fromfile(P/'selected_B_watch.uint64',dtype='<u8').reshape((len(watch),len(b)));assert len(lib)==len(b)==scan['selected_B_cells']
assert lib.min()>0 and counts.max()<=lib.max() and np.all(counts<=lib[None,:]);assert int(lib.sum())==scan['selected_B_count_sum']
expr=np.log1p(counts/lib[None,:]*10000);ifn_idx=[watch.index(g) for g in IFN];b['IFN12']=expr[ifn_idx].mean(axis=0)
# Source-coarse B average and per-source-state averages are separately retained.
agg=b.groupby('unit_id',observed=True).agg(B_IFN12=('IFN12','mean'),B_cells=('IFN12','size'))
per=per.join(agg);assert per.loc[per.B_total.gt(0),'B_total'].eq(per.loc[per.B_total.gt(0),'B_cells']).all()
per['B_fraction_all_PBMC']=per.B_total/per.n_cells
per['named_state_cells']=per[states].sum(axis=1)
per['state_annotation_assessed']=per.named_state_cells.gt(0)
per['state_label_coverage']=per.named_state_cells/per.B_total.replace(0,np.nan)
missing=b.assign(flag=b.ct_cov.isna().astype(int)).groupby('unit_id').flag.sum()
other=b.assign(flag=(b.ct_cov.notna() & ~b.ct_cov.isin(states)).astype(int)).groupby('unit_id').flag.sum()
per['B_missing_fine_label']=missing.reindex(per.index,fill_value=0)
per['B_other_labelled']=other.reindex(per.index,fill_value=0)
assert (per.B_missing_fine_label+per.B_other_labelled).eq(per.Source_B_other_or_unresolved).all()
for state in states+['Source_B_other_or_unresolved']:
 per['fraction_within_B__'+state]=per[state]/per.B_total.replace(0,np.nan)
 per['fraction_all_PBMC__'+state]=per[state]/per.n_cells
 if state in states:
  per.loc[per.B_total.gt(0)&~per.state_annotation_assessed,['fraction_within_B__'+state,'fraction_all_PBMC__'+state]]=np.nan
state_scores=b.groupby(['unit_id','state'],observed=True).agg(IFN12=('IFN12','mean'),target_cells=('IFN12','size')).reset_index()
state_scores=state_scores.merge(per.reset_index()[['unit_id','ind_cov','cohort','condition','Age','Age10','Sex','pop_cov','B_total','n_cells']],on='unit_id',validate='many_to_one');state_scores['IFN_eligible']=state_scores.target_cells.ge(20)
state_scores.to_csv(T/'B_state_IFN_by_individual.tsv',sep='\t',index=False)
per.to_csv(T/'B_individual_composition_and_IFN.tsv',sep='\t')
# Reproduce the existing aggregate B-cell score and eligibility before new models.
oldscore=pd.read_csv(OLD/'tables/GSE174188_target_cell_IFN_by_unique_donor.tsv',sep='\t');oldscore=oldscore[oldscore.cg_cov.eq('B')].set_index('ind_cov');our=per[per.B_total.gt(0)].set_index('ind_cov');joined=our.join(oldscore[['IFN12_log1p10k_mean','n_cells']].rename(columns={'n_cells':'old_B_cells'}))
max_score_error=float((joined.B_IFN12-joined.IFN12_log1p10k_mean).abs().max());assert max_score_error<1e-12
assert joined.B_total.eq(joined.old_B_cells).all()
results=[];model_inputs=[]
def fit(data,outcome,analysis,family,state,endpoint,extra=()):
 data=data.copy().reset_index(drop=False);required=['ind_cov','cohort','Age10','Sex','pop_cov','condition',outcome]+list(extra);valid=data[required].notna().all(axis=1);assert valid.all(),'Covariates/outcome missing; do not silently drop'
 assert data.ind_cov.is_unique
 data['cohort']=data.cohort.astype(int).astype(str);data['condition']=pd.Categorical(data.condition,categories=['HD','SLE']);data['y']=data[outcome]
 n_HD=int(data.condition.eq('HD').sum());n_SLE=int(data.condition.eq('SLE').sum());design='y ~ '+('C(cohort) + ' if data.cohort.nunique()>1 else '')+'Age10 + C(Sex) + C(pop_cov) + C(condition)'+''.join(' + Q("'+x+'")' for x in extra)
 out={'analysis':analysis,'family':family,'state':state,'endpoint':endpoint,'n_HD':n_HD,'n_SLE':n_SLE,'n_individuals':len(data),'median_HD':float(data.loc[data.condition.eq('HD'),'y'].median()) if n_HD else np.nan,'median_SLE':float(data.loc[data.condition.eq('SLE'),'y'].median()) if n_SLE else np.nan,'adjusted_SLE_minus_HD':np.nan,'HC3_CI_low':np.nan,'HC3_CI_high':np.nan,'HC3_SE':np.nan,'p_raw':np.nan,'status':'not_estimable','reason':'','formula':design,'covariance':'HC3','units':'fraction (multiply by 100 for percentage points)' if 'fraction' in endpoint else 'mean log1p(counts per 10,000) across IFN12 genes','n_model_columns':np.nan,'residual_df':np.nan,'max_leverage':np.nan}
 if min(n_HD,n_SLE)<10:out['reason']='Fewer than 10 eligible source individuals in at least one arm.'
 else:
  fit=smf.ols(design,data=data).fit(cov_type='HC3');key='C(condition)[T.SLE]';dm=fit.model.exog;out['n_model_columns']=dm.shape[1];out['residual_df']=float(fit.df_resid)
  if np.linalg.matrix_rank(dm)!=dm.shape[1]:out['reason']='Design matrix is not full rank.'
  elif fit.df_resid<10:out['reason']='Fewer than 10 residual degrees of freedom.'
  elif key not in fit.params or not np.isfinite(fit.bse[key]) or fit.bse[key]<=0:out['reason']='Disease coefficient/HC3 variance unavailable.'
  else:
   ci=fit.conf_int().loc[key];lev=fit.get_influence().hat_matrix_diag
   if np.max(lev)>=1-1e-9:out['reason']='Unit leverage prevents a stable HC3 estimate.'
   else:
    out.update({'status':'estimated','reason':'','adjusted_SLE_minus_HD':float(fit.params[key]),'HC3_CI_low':float(ci.iloc[0]),'HC3_CI_high':float(ci.iloc[1]),'HC3_SE':float(fit.bse[key]),'p_raw':float(fit.pvalues[key]),'max_leverage':float(np.max(lev))})
 for _,z in data.iterrows():model_inputs.append({'analysis':analysis,'family':family,'state':state,'endpoint':endpoint,'ind_cov':z.ind_cov,'condition':str(z.condition),'unit_id':z.get('unit_id',''),'y':z.y,'Age10':z.Age10,'Sex':z.Sex,'pop_cov':z.pop_cov,'cohort':z.cohort,'eligible_cell_count':z.get('target_cells',z.get('B_total',np.nan))})
 results.append(out);return out
# Same eligible individuals for B proportion and B IFN; original results are a distinct comparison.
matched=per[per.B_total.ge(20)].copy();assert matched.condition.value_counts().to_dict()=={'SLE':159,'HD':71}
fit(per,'B_fraction_all_PBMC','original_233_composition_recalculation','original_main_check','B','B_fraction_all_PBMC')
fit(matched,'B_IFN12','original_230_B_IFN_recalculation','original_main_check','B','B_IFN12')
fit(matched,'B_fraction_all_PBMC','same_eligible_B_population','matched_B_population','B','B_fraction_all_PBMC')
fit(matched,'B_IFN12','same_eligible_B_population','matched_B_population','B','B_IFN12')
for state in states:
 fit(matched[matched.state_annotation_assessed],'fraction_within_B__'+state,'source_B_state_primary','state_composition_within_B',state,'state_fraction_within_B')
 fit(per[per.state_annotation_assessed|per.B_total.eq(0)],'fraction_all_PBMC__'+state,'source_B_state_primary','state_composition_all_PBMC',state,'state_fraction_all_PBMC')
 sub=state_scores[state_scores.state.eq(state)&state_scores.IFN_eligible]
 fit(sub,'IFN12','source_B_state_primary','state_IFN',state,'state_IFN12')
eligible_states=[x['state'] for x in results if x['family']=='state_IFN' and x['status']=='estimated'];common=set(matched.ind_cov)
for state in eligible_states:common&=set(state_scores.loc[state_scores.state.eq(state)&state_scores.IFN_eligible,'ind_cov'])
common_data=per[per.ind_cov.isin(common)]
if len(eligible_states)>=2:
 for state in eligible_states:
  fit(common_data,'fraction_within_B__'+state,'common_state_eligible_population','common_state_population',state,'state_fraction_within_B')
  sub=state_scores[state_scores.state.eq(state)&state_scores.ind_cov.isin(common)]
  fit(sub,'IFN12','common_state_eligible_population','common_state_population',state,'state_IFN12')
# Additive mixture adjustment is exploratory, not an intrinsic-cell-state conclusion.
extra=['fraction_within_B__'+x for x in states]
assessed=matched[matched.state_annotation_assessed].copy()
fit(assessed,'B_IFN12','assessed_state_population_B_IFN_unadjusted','mixture_adjusted_IFN','B','B_IFN12')
fit(assessed,'B_IFN12','state_fraction_adjusted_B_IFN','mixture_adjusted_IFN','B','B_IFN12',extra=extra)
# Sensitivity cohorts reuse original selected individuals, never reselect batches.
inv=pd.read_csv(OLD/'tables/GSE174188_all_donor_batch_inventory.tsv',sep='\t');switch=set(inv.groupby('ind_cov').Status.nunique().loc[lambda v:v>1].index)
for analysis,ids in [('selected_cohort4',set(per.loc[per.cohort.eq(4),'ind_cov'])),('source_status_consistent',set(per.ind_cov)-switch)]:
 family='cohort4_state_sensitivity' if analysis=='selected_cohort4' else 'status_consistent_state_sensitivity'
 for state in states:
  fit(matched[matched.ind_cov.isin(ids)&matched.state_annotation_assessed],'fraction_within_B__'+state,analysis,family,state,'state_fraction_within_B')
  ss=state_scores[state_scores.state.eq(state)&state_scores.IFN_eligible&state_scores.ind_cov.isin(ids)]
  fit(ss,'IFN12',analysis,family,state,'state_IFN12')
# Label-coverage diagnostics requested after the initial analysis, documented separately.
for label,outcome in [('missing_fine_label','B_missing_fine_label'),('other_labelled','B_other_labelled'),('total_unresolved','Source_B_other_or_unresolved')]:
 matched['coverage_fraction__'+label]=matched[outcome]/matched.B_total
 fit(matched,'coverage_fraction__'+label,'source_label_coverage_diagnostic','source_label_coverage',label,'coverage_fraction_within_B')
high=assessed[assessed.state_label_coverage.ge(.90)]
for state in states:
 fit(high,'fraction_within_B__'+state,'high_state_annotation_coverage_90pct','annotation_coverage_sensitivity',state,'state_fraction_within_B')
 ss=state_scores[state_scores.state.eq(state)&state_scores.IFN_eligible&state_scores.ind_cov.isin(high.ind_cov)]
 fit(ss,'IFN12','high_state_annotation_coverage_90pct','annotation_coverage_sensitivity',state,'state_IFN12')
# Save coverage summaries without treating missing source labels as observed state absence.
coverage_rows=[]
for grouping in [('condition',),('cohort','condition'),('cohort','batch_cov','condition')]:
 for keys,sub in per.groupby(list(grouping),observed=True):
  keys=(keys,) if not isinstance(keys,tuple) else keys; rec=dict(zip(grouping,keys));positive=sub[sub.B_total.gt(0)];den=int(sub.B_total.sum())
  rec.update({'grouping':'_'.join(grouping),'n_individuals':len(sub),'n_ge20_B':int(sub.B_total.ge(20).sum()),'n_Bpositive_entirely_unassessed':int((sub.B_total.gt(0)&~sub.state_annotation_assessed).sum()),'source_B_cells':den,'missing_fine_label_cells':int(sub.B_missing_fine_label.sum()),'other_labelled_cells':int(sub.B_other_labelled.sum()),'unresolved_cells':int(sub.Source_B_other_or_unresolved.sum()),'cell_weighted_unresolved_fraction':float(sub.Source_B_other_or_unresolved.sum()/den) if den else np.nan,'median_individual_unresolved_fraction':float((positive.Source_B_other_or_unresolved/positive.B_total).median()),'median_individual_state_label_coverage':float(positive.state_label_coverage.median())});coverage_rows.append(rec)
coverage_summary=pd.DataFrame(coverage_rows);coverage_summary.to_csv(T/'B_source_label_coverage_by_arm_cohort_batch.tsv',sep='\t',index=False)
per.loc[per.B_total.gt(0)&~per.state_annotation_assessed].reset_index().to_csv(T/'B_unassessed_state_annotation_individuals.tsv',sep='\t',index=False)
common_data.reset_index().to_csv(T/'common_eligible_205_people.tsv',sep='\t',index=False)
r=pd.DataFrame(results);r['BH_q']=np.nan;r['planned_family_size']=0
for family,ix in r.groupby('family',sort=False).groups.items():
 if family=='original_main_check':continue # compare raw P to original; original broader FDR families unchanged
 pvals=r.loc[ix,'p_raw'].fillna(1);qs=multipletests(pvals,method='fdr_bh')[1];r.loc[ix,'BH_q']=np.where(r.loc[ix,'status'].eq('estimated'),qs,np.nan);r.loc[ix,'planned_family_size']=len(ix)
r.to_csv(T/'B_state_HC3_models.tsv',sep='\t',index=False);pd.DataFrame(model_inputs).to_csv(T/'B_state_model_input_rows.tsv',sep='\t',index=False)
# Check the two original aggregate B coefficients and raw P against their existing records.
oldmodel=pd.read_csv(OLD/'tables/GSE174188_unique_donor_HC3_composition_IFN_tests.tsv',sep='\t');original_errors=[]
for endpoint,kind in [('B_fraction_all_PBMC','composition'),('B_IFN12','IFN12')]:
 a=r[r.family.eq('original_main_check')&r.endpoint.eq(endpoint)].iloc[0];o=oldmodel[oldmodel.analysis.eq('primary_unique_batch')&oldmodel.kind.eq(kind)&oldmodel.cell_type.eq('B')].iloc[0]
 for newcol,oldcol in [('adjusted_SLE_minus_HD','adjusted_coefficient_SLE_minus_HD'),('HC3_CI_low','HC3_95CI_low'),('HC3_CI_high','HC3_95CI_high'),('p_raw','p_raw')]:
  error=abs(a[newcol]-o[oldcol]);original_errors.append({'endpoint':endpoint,'field':newcol,'absolute_error':float(error)});assert np.isclose(a[newcol],o[oldcol],atol=1e-12,rtol=1e-10)
# Marker review: donor summaries only, descriptive, source annotations are unchanged.
marker_rows=[]
for gi,g in enumerate(watch):
 if g in IFN:continue
 z=b[['unit_id','state']].copy();z['mean_log1p10k']=expr[gi];z['detection_fraction']=counts[gi].gt(0) if hasattr(counts[gi],'gt') else counts[gi]>0
 z=z.groupby(['unit_id','state'],observed=True).agg(mean_log1p10k=('mean_log1p10k','mean'),detection_fraction=('detection_fraction','mean'),target_cells=('mean_log1p10k','size')).reset_index();z['gene']=g;marker_rows.append(z)
markers=pd.concat(marker_rows,ignore_index=True).merge(per.reset_index()[['unit_id','ind_cov','condition']],on='unit_id',validate='many_to_one');markers.to_csv(T/'B_state_marker_review_by_individual.tsv',sep='\t',index=False)
ms=markers[markers.target_cells.ge(20)].groupby(['state','gene','condition'],observed=True).agg(n_individuals=('ind_cov','nunique'),median_donor_mean_log1p10k=('mean_log1p10k','median'),median_donor_detection_fraction=('detection_fraction','median')).reset_index();ms.to_csv(T/'B_state_marker_summary.tsv',sep='\t',index=False)
# Numeric and interpretation QA, no artificial rerun of original QC.
qa={'passed':True,'source_cell_order_verified':inp['matrix_metadata_barcode_order_exact'],'source_full_matrix_materialized':False,'new_cell_QC_performed':False,'selected_individuals_unique':bool(per.ind_cov.is_unique),'source_batch_selection_unchanged':True,'n_primary_original':len(per),'n_same_B_eligible':len(matched),'same_B_eligible_HD':int(matched.condition.eq('HD').sum()),'same_B_eligible_SLE':int(matched.condition.eq('SLE').sum()),'max_abs_IFN_difference_from_existing_source_score':max_score_error,'original_aggregate_model_comparisons':original_errors,'state_counts_sum_to_all_B':bool(per[states+['Source_B_other_or_unresolved']].sum(axis=1).eq(per.B_total).all()),'within_B_fractions_sum_one_when_assessed':bool(np.allclose(per.loc[per.state_annotation_assessed,['fraction_within_B__'+s for s in states+['Source_B_other_or_unresolved']]].sum(axis=1),1)),'unassessed_Bpositive_state_fraction_NA':bool(per.loc[per.B_total.gt(0)&~per.state_annotation_assessed,['fraction_within_B__'+s for s in states]].isna().all().all()),'state_composition_eligible':{'n':len(assessed),'HD':int(assessed.condition.eq('HD').sum()),'SLE':int(assessed.condition.eq('SLE').sum())},'common_population_definition':'Original selected batch and >=20 total B cells and >=20 B_naive and >=20 B_mem cells; no B_atypical requirement; same people reused across all four common-population endpoints.','eligible_state_IFN_models':eligible_states,'common_state_eligible':{'n':len(common_data),'HD':int(common_data.condition.eq('HD').sum()),'SLE':int(common_data.condition.eq('SLE').sum())},'all_estimated_models_full_rank_and_HC3':bool(r.loc[r.status.eq('estimated'),'HC3_SE'].gt(0).all()),'unavailable_models':r.loc[r.status.ne('estimated'),['analysis','state','endpoint','n_HD','n_SLE','reason']].to_dict('records'),'analysis_plan_sha256':sha(P/'analysis_plan.json'),'script_sha256':sha(Path(__file__)),'stream_cpp_sha256':sha(Path(__file__).with_name('stream_B_watch.cpp')),'counts_scan':scan,'missing_marker_features':inp['missing_requested_markers'],'independent_replication':False,'proof_of_cell_intrinsic_activation':False,'new_disease_mechanism_claim':False}
(P/'analysis_QA.json').write_text(json.dumps(qa,indent=2)+'\n')
# Editable workbook, black text and simple grey headings.
wb=Workbook();wb.remove(wb.active)
readme=pd.DataFrame([['Question','Source-recorded B-state composition and IFN12; same-eligible-person sensitivity.'],['Participants','Original selection: 162 SLE, 71 HD. Aggregate B IFN and matched composition: 159 SLE, 71 HD.'],['Units','Within-B fractions use all cg_cov=B cells; PBMC fractions use all source cells in the selected batch.'],['IFN12','Mean log1p(count/library*10,000) over 12 prespecified genes per cell; donor mean thereafter.'],['States','B_naive, B_mem, B_atypical from source ct_cov within cg_cov=B; unresolved retained separately.'],['Threshold','≥20 target-state cells for IFN; ≥10 eligible individuals per arm for an estimable state model.'],['Inference','Covariate-adjusted OLS with HC3. Individuals are observations.'],['FDR','BH by named planned family; unavailable planned tests enter adjustment as P=1, shown unavailable.'],['Interpretation','Within-source observational follow-up; sensitivity populations overlap; no independent replication or causal intrinsic-activation claim.'],['Full records','Accompanying TSVs preserve model inputs, descriptive markers and cell-ID mapping.'],['Source counts','Archived processed/remapped matrix; no new raw-droplet QC or whole-atlas reprocessing.'],['Unavailable markers',', '.join(inp['missing_requested_markers'])]],columns=['Item','Definition'])
frames={'Read_me':readme,'Coverage':pd.read_csv(T/'B_state_coverage_before_models.tsv',sep='\t'),'Label_coverage':coverage_summary,'Unassessed_state_labels':per.loc[per.B_total.gt(0)&~per.state_annotation_assessed].reset_index(),'HC3_models':r,'Individual_B_records':per.reset_index(),'State_IFN_records':state_scores,'Marker_summary':ms,'Coarse_fine_pairs':pd.read_csv(T/'source_B_coarse_fine_label_pairs.tsv',sep='\t'),'Model_input_rows':pd.DataFrame(model_inputs),'Source_files':pd.DataFrame([{'source':k,**v} for k,v in inp.items() if isinstance(v,dict) and 'sha256' in v])}
for sheet,df in frames.items():
 ws=wb.create_sheet(sheet);ws.append(list(df.columns))
 for row in df.itertuples(index=False,name=None):ws.append([None if (isinstance(v,float) and not np.isfinite(v)) else (v.item() if isinstance(v,np.generic) else v) for v in row])
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for row in ws:
  for c in row:c.font=Font(name='Arial',size=10,color='000000');c.alignment=Alignment(vertical='top',wrap_text=True)
 for c in ws[1]:c.font=Font(name='Arial',size=10,bold=True,color='000000');c.fill=PatternFill('solid',fgColor='E8E8E8')
 for i,c in enumerate(ws[1],start=1):ws.column_dimensions[get_column_letter(i)].width=42 if sheet=='Read_me' and i==2 else (28 if len(str(c.value))>15 else 18)
wb.save(W/'B_Cell_State_Analysis.xlsx')
# Candidate prose uses exact output estimates; do not insert these automatically into manuscripts.
def result(family,state,endpoint):return r[(r.family==family)&(r.state==state)&(r.endpoint==endpoint)].iloc[0]
match_comp=result('matched_B_population','B','B_fraction_all_PBMC');match_ifn=result('matched_B_population','B','B_IFN12');naive=result('state_IFN','B_naive','state_IFN12');memory=result('state_IFN','B_mem','state_IFN12');mix=r[r.analysis.eq('state_fraction_adjusted_B_IFN')].iloc[0]
fr=lambda x:f"{100*x.adjusted_SLE_minus_HD:.2f} percentage points (95% CI {100*x.HC3_CI_low:.2f} to {100*x.HC3_CI_high:.2f})"
sc=lambda x:f"{x.adjusted_SLE_minus_HD:.3f} (95% CI {x.HC3_CI_low:.3f} to {x.HC3_CI_high:.3f}; BH q={x.BH_q:.3g})"
text_en=f'''Methods candidate\nSource-recorded B_naive, B_mem and B_atypical labels were examined within the GSE174188 cg_cov=B compartment, retaining other or missing fine labels separately. The existing one-batch selection for 162 SLE and 71 healthy reference individuals was unchanged. State fractions were calculated within all B cells, retaining observed zero numerators only when state annotation was assessed. Individuals with B cells but no named-state labels were treated as unavailable for state-composition analysis. The denominator retained all B cells, including unclassified cells; models used individuals with at least 20 B cells and at least one assigned state. IFN12 was calculated from the same 12 genes and processed source-count normalization as the original analysis, then averaged within each state and individual. The sensitivity analysis restricted to at least 90% named-state coverage was specified after a source-label coverage review. State-level expression required at least 20 target cells per individual and at least 10 eligible individuals in each arm. Covariate-adjusted OLS models used age, sex, ancestry and processing cohort with HC3 covariance. BH adjustment was applied separately to the planned endpoint families. All analyses are within-source follow-up and do not provide independent replication.\n\nResults candidate\nThe selected batches contained {len(b):,} source B cells. Among 159 SLE cases and 71 references eligible for aggregate B-cell IFN analysis, the adjusted overall B-cell fraction difference was {fr(match_comp)}, while the IFN12 difference was {sc(match_ifn)} in mean log1p(counts per 10,000) units. Source-recorded naive B cells ({int(naive.n_SLE)} cases, {int(naive.n_HD)} references) and memory B cells ({int(memory.n_SLE)} cases, {int(memory.n_HD)} references) showed IFN12 differences of {sc(naive)} and {sc(memory)}, respectively. B_atypical IFN comparisons were not fitted because only three references had at least 20 such cells. Within the 155 cases and 70 references with assessed B-state annotation, the recorded naive, memory and atypical fractions differed by 8.49, −10.14 and 1.79 percentage points, respectively, using all B cells as the denominator. Five individuals with wholly unassigned fine labels were treated as unavailable for these state-composition comparisons. These estimates show how average B-cell expression and its recorded state composition can be examined separately; they do not establish cell-intrinsic activation. The matched-population results use 159 cases and 71 references; the separate composition-adjusted model uses 155 cases and 70 references with assessed source states. Both are reported alongside the original analysis rather than replacing it.\n\nInterpretation candidate\nHigher IFN12 within source-recorded naive and memory B-cell categories supports a state-level expression association beyond a change in the aggregate B-cell fraction. The state labels remain broad, and these comparisons do not remove unmeasured within-state composition, activity, treatment or technical effects. The exploratory aggregate IFN estimate after adjustment for recorded B-state fractions is {sc(mix)}; this is a model-based sensitivity analysis, not a causal decomposition.\n'''
(W/'candidate_methods_results_en.txt').write_text(text_en)
text_zh=f'''方法候选段\n在GSE174188来源cg_cov=B细胞群内，使用原ct_cov中的B_naive、B_mem和B_atypical标签，其他及缺失细标签单独保留。沿用162例SLE和71名健康参考个体的一人一批次选择。各状态比例以全部B细胞为分母，在细标签可评估时保留记录的零值；有B细胞但三种状态标签均缺失的个体，其状态组成记为不可用，不作为亚型不存在。分母保留所有B细胞，包括未细分细胞；组成模型要求每人至少20个B细胞且至少有一种状态标签。另在标签覆盖复核后制定至少90%状态注释覆盖的敏感性分析。IFN12沿用原分析的12个基因及处理后来源counts归一化，先计算细胞评分，再在每个状态及个体内取均值；状态表达比较要求每人至少20个目标细胞、每组至少10名合格个体。采用校正年龄、性别、祖源和处理队列的OLS模型及HC3协方差，并按事先记录的结局族分别进行BH校正。本分析是同一来源队列内的深化，不是独立队列验证。\n\n结果候选段\n选定批次共有{len(b):,}个来源B细胞。对同一组159例SLE和71名参考个体，总体B细胞比例的校正差异为{100*match_comp.adjusted_SLE_minus_HD:.2f}个百分点（95%置信区间{100*match_comp.HC3_CI_low:.2f}至{100*match_comp.HC3_CI_high:.2f}）；IFN12校正差异为{match_ifn.adjusted_SLE_minus_HD:.3f}（95%置信区间{match_ifn.HC3_CI_low:.3f}至{match_ifn.HC3_CI_high:.3f}），单位为平均log1p（每10,000计数）。来源记录的初始B细胞纳入{int(naive.n_SLE)}例病例和{int(naive.n_HD)}名参考个体，IFN12差异为{naive.adjusted_SLE_minus_HD:.3f}（95%置信区间{naive.HC3_CI_low:.3f}至{naive.HC3_CI_high:.3f}，BH q={naive.BH_q:.3g}）；记忆B细胞纳入{int(memory.n_SLE)}例病例和{int(memory.n_HD)}名参考个体，差异为{memory.adjusted_SLE_minus_HD:.3f}（95%置信区间{memory.HC3_CI_low:.3f}至{memory.HC3_CI_high:.3f}，BH q={memory.BH_q:.3g}）。B_atypical中仅3名参考个体达到20个细胞，因此未拟合状态内IFN组间模型。在细标签可评估的155例病例和70名参考个体中，以全部B细胞为分母，来源记录的初始、记忆及非典型B细胞比例分别相差8.49、−10.14和1.79个百分点；5名细标签完全缺失个体的状态组成记为不可用。159+71人的同人群结果与155+70人的状态比例调整结果分别与原主分析并列报告，不替换原结果，也不据此声称细胞内在激活。\n\n讨论候选段\n在来源记录的初始及记忆B细胞中观察IFN表达差异，有助于进一步区分细胞总体比例与状态内平均表达。现有标签仍较粗，不能排除状态内部组成、疾病活动、治疗或其他技术因素。加入已记录状态比例后，在155例病例和70名参考个体这一相同可评估状态人群中，探索性总体B细胞IFN效应为{mix.adjusted_SLE_minus_HD:.3f}（95%置信区间{mix.HC3_CI_low:.3f}至{mix.HC3_CI_high:.3f}）；这属于模型敏感性分析，不是因果分解。\n'''
(W/'candidate_methods_results_zh.txt').write_text(text_zh)
not_done=[{'item':'Entirely unassessed source B-state labels','reason':'Five B-positive individuals (4 SLE and 1 HD, processing cohort 3) have no named B-state labels. Their state-composition outcomes are NA, not zeros; they remain in aggregate B-cell analyses. Partial unresolved labels remain in all-B denominators and are quantified by condition/cohort/batch.'},{'item':'B_atypical within-state IFN comparison','reason':'Only 3 HD individuals satisfy >=20 source B_atypical cells; the prespecified minimum is 10 in each arm. No model fitted and no relaxed threshold substituted.'},{'item':'Independent B-state cohort replication','reason':'All primary and sensitivity participants come from GSE174188 and overlap; no independent verified cohort is substituted.'},{'item':'Fine B-subtype biological validation','reason':'Source ct_cov labels are broad; available marker panels are reviewed descriptively, not a blinded independent fine-annotation benchmark.'},{'item':'IGHD, IGHM and JCHAIN marker checks','reason':'These requested markers are absent from the retained source features; no counts imputed or unrelated aliases substituted.'},{'item':'Clinical activity/treatment adjustment and causal intrinsic-activation inference','reason':'Source Status is not a harmonized quantitative activity/treatment covariate and does not resolve confounding. No new mechanism or intrinsic-activation claim.'}]
(W/'not_completed.json').write_text(json.dumps(not_done,indent=2)+'\n')
print(r[['analysis','family','state','endpoint','n_HD','n_SLE','adjusted_SLE_minus_HD','HC3_CI_low','HC3_CI_high','BH_q','status']].to_string(index=False),flush=True)
