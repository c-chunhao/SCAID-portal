from pathlib import Path
import os
import pandas as pd,numpy as np,json,hashlib,shutil,sys
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
W=Path(os.environ.get('SCAID_DIAGNOSTIC_DIR',str(Path(__file__).resolve().parent)))
B=Path(os.environ.get('SCAID_MARKER_EVIDENCE_DIR','marker_evidence'))
I=Path(os.environ.get('SCAID_INTEGRATION_REVIEW_DIR','integration_annotation'))
def sha(p):
 h=hashlib.sha256()
 with open(p,'rb') as f:
  for v in iter(lambda:f.read(16*1024*1024),b''):h.update(v)
 return h.hexdigest()
def js(p,x):Path(p).write_text(json.dumps(x,ensure_ascii=False,indent=2,allow_nan=False)+'\n')
def df(fn):return pd.read_csv(B/fn,sep='\t')
cov=df('all_63_fine_label_audit_coverage.tsv');review=df('source_fine_label_coherence_review.tsv.gz');cand=df('source_review_candidates.tsv');joint=df('actual_joint/actual_joint_source_evidence.tsv');bounds=df('actual_joint/135_bound_candidates_actual_joint_review.tsv');genes=df('full_RNA_marker_gene_coverage.tsv');dc=pd.read_csv(I/'T1D_DC_Mono_candidate_review.tsv.gz',sep='\t')
old=json.loads((B/'marker_review_summary.json').read_text());oldj=json.loads((B/'actual_joint/actual_joint_validation_summary.json').read_text())
assert len(cov)==63 and cov.Cells.sum()==1191460
assert cov.groupby('Lineage').Cells.sum().to_dict()=={'B_Plasma':253635,'T_NK':937825}
assert len(review)==2326 and review.Cells.sum()==1191460 and review.Review_min20_cells.sum()==1800
qual_sources=review[review.Review_min20_cells].groupby(['Lineage','celltype2']).Dataset.nunique()
assert len(qual_sources)==63 and int(qual_sources.min())==3
cov['Distinct_source_IDs_ge20']=[int(qual_sources.loc[(r.Lineage,r.celltype2)]) for r in cov.itertuples()]
assert len(cand)==259 and cand.Review_min20_cells.all()
assert len(joint)==443 and joint.Cells.sum()==467302 and not joint.duplicated(['Data_object','Lineage','celltype2']).any()
assert len(bounds)==135 and int(bounds.gamma_program_joint_Fraction.ge(.10).sum())==120 and int(bounds.gamma_program_strict_joint_Fraction.ge(.10).sum())==101
pl=joint[joint.celltype2.eq('Plasmablast')]
av=pl.plasma_2of3_joint_Available
assert pl.Cells.sum()==755 and pl.loc[av,'Cells'].sum()==325 and pl.loc[av,'plasma_2of3_joint_Detected_cells'].sum()==222
assert pl.loc[~av,'Cells'].sum()==430
assert pl.loc[pl.MZB1_SDC1_actual_joint_Available,'Cells'].sum()==745
assert pl.loc[pl.MZB1_SDC1_actual_joint_Available,'MZB1_SDC1_actual_joint_Detected_cells'].sum()==188
for c in joint:
 if c.endswith('_Detected_cells'):
  prefix=c[:-len('_Detected_cells')];fc=prefix+'_Fraction';ac=prefix+'_Available'
  if fc in joint and ac in joint:
   v=joint[ac].astype(bool)
   assert (joint.loc[v,c]>=0).all() and (joint.loc[v,c]<=joint.loc[v,'Cells']).all()
   assert np.allclose(joint.loc[v,c]/joint.loc[v,'Cells'],joint.loc[v,fc],atol=1e-12)
   assert joint.loc[~v,fc].isna().all()
dcsub=dc[dc.source_author_coarse.eq('DC')&dc.legacy_curated_coarse.eq('Mononuclear Phagocytes')]
assert len(dcsub)==324
flags=Counter()
for item in cand.Review_flags:
 for f in str(item).split(';'):flags[f]+=1
assert dict(flags)==old['flags']
# Concise results for Word; full rows are supplied in the workbook and TSV companions.
summary_rows=[
 ['Fine labels reviewed','63','43 T/NK and 20 B/plasma'],
 ['Cells represented','1,191,460','937,825 T/NK; 253,635 B/plasma'],
 ['Source-object × fine-label groups','2,326','1,800 groups have at least 20 cells'],
 ['Specimen × fine-label groups','25,738','11,583 groups have at least 20 cells'],
 ['Source groups flagged for review','259','45 fine labels across 30 source objects; flags may overlap'],
 ['Cells in targeted co-detection review','467,302','443 source-object × fine-label groups; 46 source objects']]
summary_zh=[
 ['已检查细标签','63','T/NK 43类；B/浆细胞20类'],
 ['所覆盖细胞','1,191,460','T/NK 937,825；B/浆细胞253,635'],
 ['来源对象×细标签组合','2,326','其中1,800组至少20个细胞'],
 ['标本×细标签组合','25,738','其中11,583组至少20个细胞'],
 ['提示进一步复核的来源组合','259','涉及45个细标签、30个来源对象；提示可重叠'],
 ['定向共检出检查细胞','467,302','443个来源对象×细标签组合；46个来源对象']]
risk_rows=[
 ['T/CD3–TRD RNA programme','135 initial source-group candidates','120/135 meet the measured ≥10% joint-detection screen; 101/135 meet the stricter all-three-CD3 screen','RNA programme evidence; not confirmed TCR identity or doublet exclusion'],
 ['NKT_KLRC2: GSE216883','745 cells','T/CD3–NK joint detection 12.75%; all-three-CD3 detection 5.77%','Review the NK-like versus NKT interpretation'],
 ['NKT_KLRC2: GSE157278','1,503 cells','T/CD3–NK joint detection 6.45%; all-three-CD3 detection 1.06%','Review the NK-like versus NKT interpretation'],
 ['Bgc_MKI67−: GSE174188 / GSE221297','21 / 32 cells','Low BCL6/AICDA detection in the source-specific screen','Small groups and dropout require marker/context review'],
 ['Plasmablast','755 cells in 25 source objects','Complete three-gene panel measurable in 325 cells; 222 co-detect ≥2 plasma genes; 430 lack required features','Unavailable features are not negative expression; do not report 222/755 as accuracy'],
 ['Source cDC assigned to mononuclear phagocytes: syn53641849','324 cells','298 monocyte-marker support; 22 mixed; 4 sparse/non-specific under the existing panel','Targeted expert review; neither source nor curated label is automatically treated as truth']]
risk_zh=[
 ['T/CD3–TRD RNA程序','135个初筛来源组合','120/135达到实际联合检出≥10%标准；101/135达到三个CD3均检出的更严格标准','属于RNA程序依据，不能确认TCR身份或排除双细胞'],
 ['NKT_KLRC2：GSE216883','745个细胞','T/CD3与NK程序共检出12.75%；三个CD3均检出5.77%','优先复核NK样与NKT解释'],
 ['NKT_KLRC2：GSE157278','1,503个细胞','T/CD3与NK程序共检出6.45%；三个CD3均检出1.06%','优先复核NK样与NKT解释'],
 ['Bgc_MKI67−：GSE174188 / GSE221297','21 / 32个细胞','来源分层初筛中BCL6/AICDA检出较低','小规模及稀疏检出需结合标志基因与背景审阅'],
 ['浆母细胞','25个来源对象的755个细胞','325个细胞可测完整三基因程序，其中222个共检出至少2个浆细胞基因；430个缺必需feature','缺失基因不是表达阴性；222/755不能解释为准确率'],
 ['原作者cDC被归为单核吞噬细胞：syn53641849','324个细胞','现有面板下298个有单核细胞标志支持、22个混合、4个稀疏/非特异','需专家定向复核，不将来源或现有标签自动作为真值']]
foot_en=['Review flags identify source-conditioned marker patterns or missing marker coverage, not confirmed annotation errors. No production labels, counts or cells were changed.','The targeted second pass measures actual within-cell co-detection. Its 467,302 cells are a selected subset, not a second complete scan of all 63 labels. Source-object and specimen labels do not imply independent studies or donors.','Detection means an original count greater than zero. Unavailable required genes yield an unavailable programme. JCHAIN uses an explicitly recorded IGJ fallback only when JCHAIN is absent.']
foot_zh=['复核提示反映特定来源的标志基因模式或面板缺失，不等于确认注释错误。本次未修改生产标签、计数或细胞。','定向第二轮测量同一细胞的实际共检出，467,302个细胞是选定子集，并非再次扫描全部63类的所有细胞。来源对象及标本标签不代表独立研究或独立供体。','检出定义为原始计数大于零。缺少必需基因时程序记为不可测；仅在JCHAIN缺失时按已记录规则使用旧名称IGJ。']
tables=[{'id':'annotation_review_coverage','title_en':'Source-conditioned review of fine atlas labels','title_zh':'细分图谱标签的来源分层复核','headers_en':['Measure','Result','Scope'],'headers_zh':['指标','结果','范围'],'rows_en':summary_rows,'rows_zh':summary_zh,'footnotes_en':foot_en,'footnotes_zh':foot_zh},
 {'id':'priority_annotation_review','title_en':'Priority annotation review evidence','title_zh':'重点注释复核依据','headers_en':['Population or signal','Denominator','Observed evidence','Interpretation'],'headers_zh':['细胞群或信号','分母','观察依据','解释'],'rows_en':risk_rows,'rows_zh':risk_zh,'footnotes_en':foot_en,'footnotes_zh':foot_zh}]
res={'tables':tables,'source_flag_counts':dict(flags),'no_production_changes':True,'results_en':'Source-conditioned marker review covered all 1,191,460 atlas cells and all 63 fine labels, comprising 2,326 source-object–label groups. All fine labels were represented by groups of at least 20 cells in at least three distinct source identifiers. Of 1,800 groups with at least 20 cells, 259 had marker-pattern or feature-coverage flags for review. A targeted analysis measured same-cell marker co-detection in 467,302 selected cells. It prioritized NK/NKT interpretation in two NKT_KLRC2 groups and identified incomplete plasma-marker coverage in 430 of the 755 plasmablast cells. These checks guide source-specific expert review rather than estimate independent annotation accuracy.',
'results_zh':'来源分层标志基因检查覆盖全部1,191,460个图谱细胞和63个细标签，共2,326个来源对象—标签组合。所有细标签均在至少三个不同来源标识符中具有至少20个细胞的组合。在至少20个细胞的1,800组中，259组因标志基因模式或面板可用性被提示进一步复核。定向分析对467,302个选定细胞测量实际同细胞共检出，指出两个NKT_KLRC2来源组合需优先审查NK/NKT解释，并发现755个浆母细胞中430个缺少必需的完整浆细胞面板。这些检查用于来源特异的专家复核，不作为独立注释准确率。',
'methods_en':'Existing full-RNA marker summaries were reconciled by source object, lineage and fine label. Coverage and cell totals were checked against the saved atlas metadata. Groups with fewer than 20 cells were retained but not used in the review screens. The targeted co-detection results were checked for valid available-gene denominators, detection-count/fraction agreement and consistency with the preceding marginal bounds. RNA-programme thresholds were used only to prioritize expert review; no labels were reassigned. Source-author DC/mononuclear-phagocyte discrepancies were summarized separately using the existing cell-level marker review.',
'methods_zh':'按来源对象、谱系及细标签核对既有完整RNA标志基因汇总，并与图谱元数据核对覆盖及细胞总量。少于20个细胞的组合保留记录但不用于复核筛查。对定向共检出结果检查可用基因分母、检出细胞数与比例一致性，以及其与前轮边际上下界的关系。RNA程序阈值仅用于安排专家复核，不据此重标注。原作者DC与单核吞噬细胞分歧采用既有逐细胞标志基因结果单独汇总。'}
js(W/'annotation_review_tables.json',res)
# Copy concise original evidence, never expression matrices or giant Seurat objects.
evd=W/'annotation_evidence';evd.mkdir(exist_ok=True)
files=[B/'all_source_fine_label_marker_evidence.tsv.gz',B/'all_specimen_fine_label_marker_evidence.tsv.gz',B/'actual_joint/actual_joint_specimen_evidence.tsv.gz',B/'object_processing_QA.tsv',B/'legacy_marker_symbol_alias_availability.tsv',B/'all_63_fine_label_audit_coverage.tsv',B/'source_fine_label_coherence_review.tsv.gz',B/'source_review_candidates.tsv',B/'full_RNA_marker_gene_coverage.tsv',B/'predeclared_marker_panel.json',B/'predeclared_marker_review_rules.json',B/'marker_coherence_QA.json',B/'marker_review_summary.json',B/'actual_joint/actual_joint_source_evidence.tsv',B/'actual_joint/135_bound_candidates_actual_joint_review.tsv',B/'actual_joint/actual_joint_vs_marginal_bounds_regression.tsv',B/'actual_joint/predeclared_actual_joint_panel.json',B/'actual_joint/actual_joint_QA.json',B/'actual_joint/actual_joint_validation_summary.json',B/'audit_full_RNA_marker_coherence.py',B/'build_marker_review_evidence.py',B/'audit_targeted_actual_joint_markers.py',I/'T1D_DC_Mono_candidate_review.tsv.gz',I/'T1D_source_DC_marker_and_candidate_review_counts.tsv']
provenance=[]
for p in files:
 dest=evd/p.name;shutil.copy2(p,dest);provenance.append({'source':str(p),'source_sha256':sha(p),'copy':str(dest),'copy_sha256':sha(dest),'bytes':dest.stat().st_size})
assert all(p['source_sha256']==p['copy_sha256'] for p in provenance)
# Editable workbook; raw unavailable values are blank, with explicit available columns preserved.
book=Workbook();book.remove(book.active)
frames={
 'Read_me':pd.DataFrame([['Scope','Existing marker evidence reorganized and independently reconciled; no annotation replacement'],['Labels','43 T/NK and 20 B/plasma'],['Interpretation','Review flags are not an accuracy measure'],['Selected second pass','467,302 selected cells, not all atlas cells'],['Source unit','Source object × fine label; not independent studies or donors']],columns=['Item','Description']),
 'Review_coverage':pd.DataFrame(summary_rows,columns=['Measure','Result','Scope']),
 'Priority_review':pd.DataFrame(risk_rows,columns=['Population or signal','Denominator','Observed evidence','Interpretation']),
 'All_63_labels':cov,'Source_label_review':review,'Review_candidates':cand,'Actual_joint_443':joint,'TRD_candidates_135':bounds,'Marker_availability':genes,'DC_Mono_324':dcsub,
 'Source_files':pd.DataFrame(provenance)}
for title,data in frames.items():
 ws=book.create_sheet(title);ws.append(list(data.columns))
 for rr in data.itertuples(index=False,name=None):ws.append([None if pd.isna(v) else v.item() if isinstance(v,np.generic) else v for v in rr])
 for c in ws[1]:c.font=Font(name='Arial',size=10,bold=True,color='000000');c.fill=PatternFill('solid',fgColor='EAEAEA');c.alignment=Alignment(wrap_text=True,vertical='top');c.border=Border(top=Side(style='thin',color='000000'),bottom=Side(style='thin',color='000000'))
 for row in ws.iter_rows(min_row=2):
  for c in row:c.font=Font(name='Arial',size=10,color='000000');c.alignment=Alignment(wrap_text=True,vertical='top')
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for col in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(col)].width=45 if title in ['Read_me','Priority_review','Source_files'] else 23
 ws.row_dimensions[1].height=42;ws.sheet_view.showGridLines=False
book.save(W/'Atlas_Annotation_Review_Evidence.xlsx')
qa={'passed':True,'original_atlas_cells_checked':1191460,'fine_labels':63,'minimum_distinct_source_IDs_ge20_per_fine_label':3,'source_label_groups':2326,'source_groups_ge20':1800,'review_candidates':259,'candidate_labels':int(cand.celltype2.nunique()),'candidate_source_objects':int(cand.Data_object.nunique()),'actual_joint_cells':467302,'actual_joint_groups':443,'all_detection_counts_and_fractions_reconciled':True,'all_unavailable_program_fractions_remain_missing':True,'plasmablast_three_gene_available_cells':325,'plasmablast_positive_cells':222,'plasmablast_missing_required_features':430,'DC_discrepancy_cells':324,'DC_marker_categories':{str(k):int(v) for k,v in dcsub.DC_Mono_marker_review.value_counts().items()},'source_copies':provenance,'outputs':{'xlsx_sha256':sha(W/'Atlas_Annotation_Review_Evidence.xlsx'),'word_tables_json_sha256':sha(W/'annotation_review_tables.json')},'versions':{'python':sys.version,'pandas':pd.__version__,'numpy':np.__version__}}
js(W/'annotation_summary_QA.json',qa)
print(json.dumps({k:v for k,v in qa.items() if k not in ['source_copies','versions']},ensure_ascii=False,indent=2))
