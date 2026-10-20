"""Additional GSE174188 source-pool sensitivity using existing individual summaries.

Never reselect participants, read counts, or overwrite the original HC3 analyses.
Public reruns may supply equivalent individual-level summary inputs through environment variables.
"""
from pathlib import Path
import datetime, hashlib, json, os
import numpy as np
import pandas as pd
import scipy.stats as st
import statsmodels.formula.api as smf
from statsmodels.stats.multitest import multipletests

OUT=Path(os.environ.get('SCAID_POOL_OUTPUT','work/source_pool_sensitivity')).resolve()
OLD=Path(os.environ.get('SCAID_BASELINE_SCIENCE','data/baseline_science')).resolve()
BCELL=Path(os.environ.get('SCAID_BCELL_OUTPUT','work/bcell_state')).resolve()
for name in ['tables','provenance']: (OUT/name).mkdir(parents=True,exist_ok=True)
T=OUT/'tables'; P=OUT/'provenance'
sha=lambda p: hashlib.sha256(Path(p).read_bytes()).hexdigest()
inputs={
 'composition':OLD/'tables/GSE174188_unique_donor_composition_complete_zeros.tsv',
 'IFN':OLD/'tables/GSE174188_target_cell_IFN_by_unique_donor.tsv',
 'inventory':OLD/'tables/GSE174188_all_donor_batch_inventory.tsv',
 'original_models':OLD/'tables/GSE174188_unique_donor_HC3_composition_IFN_tests.tsv',
 'B_individual':BCELL/'tables/B_individual_composition_and_IFN.tsv',
 'B_model_rows':BCELL/'tables/B_state_model_input_rows.tsv',
 'B_models':BCELL/'tables/B_state_HC3_models.tsv'}
plan={'recorded_at_UTC':datetime.datetime.now(datetime.timezone.utc).isoformat(),
 'type':'Additional sensitivity after review of shared source processing pools; not an original prespecified primary analysis.',
 'population':'Keep original one-selected-batch per individual and each existing endpoint eligibility; >=20 target cells for IFN and >=10 people per arm.',
 'variants':['Original age/sex/ancestry/cohort design with CR1 by source batch_cov and t(G-1) inference.', 'Age/sex/ancestry/source batch_cov fixed effects with CR1 by batch_cov and t(G-1) inference; processing cohort omitted because nested.'],
 'CR1_correction':'G/(G-1)*(N-1)/(N-K), where K is full-rank design width.',
 'multiplicity':'BH separately within each pre-existing endpoint family and sensitivity variant; unavailable planned state IFN tests enter as P=1.',
 'validation':'Independent manual grouped-score covariance for every estimable model and leave-one-pool-out point estimates for key endpoints.',
 'limits':'Approximately 20 clusters; t intervals are finite-cluster approximations, not exact or unbiased guarantees. Observational within-source associations, not independent replication.',
 'source_counts_read':False,'original_results_overwritten':False}
(P/'analysis_plan.json').write_text(json.dumps(plan,indent=2)+'\n')
(P/'inputs.json').write_text(json.dumps({k:{'path':str(v),'sha256':sha(v)} for k,v in inputs.items()},indent=2)+'\n')
d={k:pd.read_csv(v,sep='\t') for k,v in inputs.items()}
comp=d['composition']; bp=d['B_individual']; inv=d['inventory']
assert comp.ind_cov.is_unique and bp.ind_cov.is_unique
assert inv.groupby('batch_cov').cohort.nunique().max()==1
meta=comp[['ind_cov','batch_cov','cohort','condition']]
specs=[]
for state in ['T4','T8','B','NK','cM','ncM','cDC','pDC']:
 z=comp.copy();z['y']=z[state]
 ref=d['original_models'].query('analysis=="primary_unique_batch" and kind=="composition" and cell_type==@state').iloc[0]
 specs.append(dict(key='original_fraction_'+state, family='original_8_PBMC_fractions',state=state,endpoint='PBMC_fraction',analysis='original_233',data=z,extra=[],reference=ref.to_dict(),reference_kind='original'))
for state in ['cM','T4','B']:
 z=d['IFN'].query('cg_cov==@state and n_cells>=20').merge(comp[['ind_cov','batch_cov','Age10','Sex','pop_cov']],on='ind_cov',validate='one_to_one');z['y']=z.IFN12_log1p10k_mean
 ref=d['original_models'].query('analysis=="primary_unique_batch" and kind=="IFN12" and cell_type==@state').iloc[0]
 specs.append(dict(key='original_IFN_'+state,family='original_3_IFN',state=state,endpoint='IFN12',analysis='original_IFN',data=z,extra=[],reference=ref.to_dict(),reference_kind='original'))
families=['matched_B_population','state_composition_within_B','state_IFN','common_state_population','mixture_adjusted_IFN','source_label_coverage']
for _,row in d['B_models'][d['B_models'].family.isin(families)].iterrows():
 q=d['B_model_rows'];z=q[q.analysis.eq(row.analysis)&q.family.eq(row.family)&q.state.eq(row.state)&q.endpoint.eq(row.endpoint)].copy()
 z=z.merge(bp[['ind_cov','batch_cov']+(['fraction_within_B__B_naive','fraction_within_B__B_mem','fraction_within_B__B_atypical'] if row.analysis=='state_fraction_adjusted_B_IFN' else [])],on='ind_cov',validate='one_to_one')
 extra=['fraction_within_B__'+s for s in ['B_naive','B_mem','B_atypical']] if row.analysis=='state_fraction_adjusted_B_IFN' else []
 specs.append(dict(key='B_'+row.analysis+'_'+row.state+'_'+row.endpoint,family=row.family,state=row.state,endpoint=row.endpoint,analysis=row.analysis,data=z,extra=extra,reference=row.to_dict(),reference_kind='B'))
assert len(specs)==28
results=[]; private_rows=[];validation=[];lobo=[];coverage=[]
for spec in specs:
 z=spec['data'].copy().reset_index(drop=True)
 assert z.ind_cov.is_unique and z[['y','Age10','Sex','pop_cov','cohort','condition','batch_cov']+spec['extra']].notna().all().all()
 z['cohort']=z.cohort.astype(str);z['condition']=pd.Categorical(z.condition,categories=['HD','SLE'])
 groups=z.groupby('batch_cov',observed=True)
 mixed=groups.condition.nunique().eq(2)
 group_meta={'N':len(z),'n_SLE':int(z.condition.eq('SLE').sum()),'n_HD':int(z.condition.eq('HD').sum()),'n_pools':z.batch_cov.nunique(),'n_mixed_condition_pools':int(mixed.sum()),'N_in_mixed_pools':int(z.batch_cov.isin(mixed[mixed].index).sum()),'n_singleton_pools':int(groups.size().eq(1).sum()),'min_pool_size':int(groups.size().min()),'max_pool_size':int(groups.size().max())}
 coverage.append(dict(model_key=spec['key'],family=spec['family'],**group_meta))
 priv=z.assign(model_key=spec['key']);private_rows.append(priv[['model_key','ind_cov','batch_cov','cohort','condition','y','Age10','Sex','pop_cov']+spec['extra']])
 for variant in ['cohort_design_CR1','pool_fixed_effects_CR1']:
  terms=('C(cohort)' if variant=='cohort_design_CR1' else 'C(batch_cov)')+' + Age10 + C(Sex) + C(pop_cov) + C(condition)'
  formula='y ~ '+terms+''.join(' + Q("'+x+'")' for x in spec['extra'])
  r={'model_key':spec['key'],'analysis':spec['analysis'],'family':spec['family'],'state':spec['state'],'endpoint':spec['endpoint'],'variant':variant,**group_meta,'formula':formula,'status':'not_estimable','reason':'','coefficient':np.nan,'CR1_SE':np.nan,'CI_low':np.nan,'CI_high':np.nan,'p_raw':np.nan,'BH_q':np.nan,'K':np.nan,'rank':np.nan,'df_t':group_meta['n_pools']-1,'CR1_factor':np.nan,'max_leverage':np.nan,'original_HC3_coefficient':np.nan,'original_HC3_CI_low':np.nan,'original_HC3_CI_high':np.nan,'original_HC3_q':np.nan}
  old=spec['reference']
  if spec['reference_kind']=='original':
   r.update(original_HC3_coefficient=old['adjusted_coefficient_SLE_minus_HD'],original_HC3_CI_low=old['HC3_95CI_low'],original_HC3_CI_high=old['HC3_95CI_high'],original_HC3_q=old['p_BH'])
  else:r.update(original_HC3_coefficient=old['adjusted_SLE_minus_HD'],original_HC3_CI_low=old['HC3_CI_low'],original_HC3_CI_high=old['HC3_CI_high'],original_HC3_q=old['BH_q'])
  if min(group_meta['n_SLE'],group_meta['n_HD'])<10:r['reason']='Fewer than 10 eligible individuals in at least one condition; original threshold retained.'
  elif group_meta['n_pools']<2:r['reason']='Fewer than two source pools.'
  else:
   f=smf.ols(formula,z).fit();X=f.model.exog;N,K=X.shape;rank=np.linalg.matrix_rank(X);r.update(K=K,rank=rank,max_leverage=float(f.get_influence().hat_matrix_diag.max()))
   if rank!=K:r['reason']='Rank-deficient design; no pseudoinverse disease inference substituted.'
   elif f.df_resid<10:r['reason']='Fewer than 10 residual degrees of freedom.'
   else:
    names=f.model.exog_names;idx=names.index('C(condition)[T.SLE]')
    robust=f.get_robustcov_results(cov_type='cluster',groups=z.batch_cov,use_correction=True,df_correction=True,use_t=True)
    G=group_meta['n_pools'];factor=G/(G-1)*(N-1)/(N-K);bread=np.linalg.inv(X.T@X);meat=np.zeros((K,K));res=np.asarray(f.resid)
    for _,ii in z.groupby('batch_cov').groups.items():
     score=X[np.asarray(ii)].T@res[np.asarray(ii)];meat+=np.outer(score,score)
    manual=bread@meat@bread*factor;err=float(np.max(np.abs(manual-robust.cov_params())))
    assert np.allclose(manual,robust.cov_params(),rtol=1e-8,atol=1e-11)
    se=float(np.sqrt(manual[idx,idx]));estimate=float(f.params.iloc[idx]);crit=float(st.t.ppf(.975,G-1));pv=float(2*st.t.sf(abs(estimate/se),G-1));ci=robust.conf_int()[idx]
    assert np.allclose(ci,[estimate-crit*se,estimate+crit*se],atol=1e-10)
    assert np.isclose(pv,robust.pvalues[idx],atol=1e-12)
    if variant=='cohort_design_CR1':assert np.isclose(estimate,r['original_HC3_coefficient'],atol=1e-11),spec['key']
    r.update(status='estimated',coefficient=estimate,CR1_SE=se,CI_low=float(ci[0]),CI_high=float(ci[1]),p_raw=pv,CR1_factor=factor)
    validation.append({'model_key':spec['key'],'variant':variant,'max_covariance_abs_error':err,'coefficient_same_as_original':True if variant=='cohort_design_CR1' else None,'df_t':G-1})
    # Leave one pool out checks point-estimate sensitivity, not new independent replications.
    if spec['key'] in ['original_fraction_cM','original_IFN_B'] or spec['family'] in ['state_IFN','common_state_population'] and spec['endpoint'].endswith('IFN12'):
     for pool in sorted(z.batch_cov.unique()):
      zz=z[z.batch_cov.ne(pool)].copy();ff=smf.ols(formula,zz).fit();XX=ff.model.exog
      rr={'model_key':spec['key'],'variant':variant,'omitted_pool':pool,'N':len(zz),'rank':int(np.linalg.matrix_rank(XX)),'K':XX.shape[1],'coefficient':np.nan,'status':'rank_deficient'}
      if rr['rank']==rr['K']:rr.update(coefficient=float(ff.params['C(condition)[T.SLE]']),status='estimated')
      lobo.append(rr)
  results.append(r)
r=pd.DataFrame(results)
for (variant,family),ii in r.groupby(['variant','family']).groups.items():
 p=r.loc[ii,'p_raw'].fillna(1);q=multipletests(p,method='fdr_bh')[1];r.loc[ii,'BH_q']=np.where(r.loc[ii,'status'].eq('estimated'),q,np.nan);r.loc[ii,'planned_family_size']=len(ii)
r.to_csv(T/'source_pool_sensitivity_models.tsv',sep='\t',index=False)
pd.DataFrame(coverage).to_csv(T/'source_pool_eligibility.tsv',sep='\t',index=False)
pd.concat(private_rows,ignore_index=True).to_csv(T/'PRIVATE_model_input_rows.tsv',sep='\t',index=False)
pd.DataFrame(lobo).to_csv(T/'PRIVATE_leave_one_pool_out.tsv',sep='\t',index=False)
lo=pd.DataFrame(lobo);los=lo.groupby(['model_key','variant']).agg(n_omissions=('omitted_pool','size'),n_estimable=('coefficient','count'),coefficient_min=('coefficient','min'),coefficient_max=('coefficient','max')).reset_index()
los.to_csv(T/'leave_one_pool_out_summary.tsv',sep='\t',index=False)
cluster_cohort=inv.groupby('batch_cov').agg(cohort=('cohort','first'),n_individual_batches=('ind_cov','size'),n_individuals=('ind_cov','nunique'))
semantics={'source_paper':'https://pmc.ncbi.nlm.nih.gov/articles/PMC9297655/', 'source_paper_doi':'10.1126/science.abf1970','source_author_code_doi':'10.5281/zenodo.4724043','source_author_code_file':'Figure_1_2.Final.ipynb','source_author_code_evidence':'Author code uses batch_cov to select the multiplexed label dmx_YS-JY-22_pool6 and separately plots Processing_Cohort (cells 7, 8 and 20).','paper_processing':'23 multiplexed pools profiled across four processing batches; donor identity recovered by genotype demultiplexing.', 'local_field_mapping':'batch_cov identifies source multiplexed pools and is nested in the four Processing_Cohort values (inventory field cohort); 20 pools are retained by the unchanged one-source-batch-per-individual selection. The mapping is supported jointly by author code, source metadata and the primary paper.', 'all_inventory_pools':len(cluster_cohort),'all_inventory_cohorts':int(inv.cohort.nunique()),'pools_per_cohort':{str(k):int(v) for k,v in cluster_cohort.groupby('cohort').size().items()},'pools_completely_nested_in_cohort':bool(inv.groupby('batch_cov').cohort.nunique().eq(1).all()),'selected_primary_pools':int(comp.batch_cov.nunique())}
(P/'source_pool_semantics.json').write_text(json.dumps(semantics,indent=2)+'\n')
qa={'passed':True,'n_endpoint_specs':len(specs),'n_model_rows':len(r),'estimated':int(r.status.eq('estimated').sum()),'unavailable':r.loc[r.status.ne('estimated'),['model_key','variant','n_SLE','n_HD','reason']].to_dict('records'),'manual_covariance_checks':len(validation),'max_manual_covariance_abs_error':max(x['max_covariance_abs_error'] for x in validation),'all_cohort_design_coefficients_preserved':True,'no_cell_matrix_read':True,'no_new_participant_or_cell_filters':True,'cohort_removed_from_pool_FE':True,'FDR_separate_by_design_and_existing_family':True,'t_df':'G-1','small_cluster_inference_is_approximate':True,'input_files_unchanged':{k:sha(v)==json.loads((P/'inputs.json').read_text())[k]['sha256'] for k,v in inputs.items()},'script_sha256':sha(Path(__file__)),'analysis_plan_sha256':sha(P/'analysis_plan.json')}
assert all(qa['input_files_unchanged'].values())
(P/'numeric_QA.json').write_text(json.dumps(qa,indent=2)+'\n')
(P/'manual_covariance_checks.json').write_text(json.dumps(validation,indent=2)+'\n')
clean=lambda df:json.loads(df.to_json(orient='records'))
(OUT/'source_pool_results.json').write_text(json.dumps({'plan':plan,'semantics':semantics,'models':clean(r),'LOBO':clean(los),'QA':qa},indent=2)+'\n')
with pd.ExcelWriter(OUT/'Source_Pool_Sensitivity.xlsx',engine='openpyxl') as x:
 pd.DataFrame([{'Item':k,'Definition':json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else str(v)} for k,v in plan.items()]).to_excel(x,sheet_name='Read_me',index=False)
 r.to_excel(x,sheet_name='Pool_sensitivity_models',index=False);pd.DataFrame(coverage).to_excel(x,sheet_name='Eligible_populations',index=False);los.to_excel(x,sheet_name='Leave_one_pool_out_summary',index=False)
 for ws in x.book:
  ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
  from openpyxl.styles import Font,Alignment
  for row in ws:
   for c in row:c.font=Font(name='Arial',size=10,color='000000');c.alignment=Alignment(vertical='top',wrap_text=True)
  for c in ws[1]:c.font=Font(name='Arial',size=10,bold=True,color='000000')
  from openpyxl.utils import get_column_letter
  for i in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(i)].width=28
print(r[['model_key','variant','N','coefficient','CI_low','CI_high','BH_q','status']].to_string(index=False),flush=True)
